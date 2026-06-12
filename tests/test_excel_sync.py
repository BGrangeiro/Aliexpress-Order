from __future__ import annotations

from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

import aliexpress_orders_sync.excel_sync as excel_sync
from aliexpress_orders_sync.excel_sync import sync_orders_to_excel, sync_orders_to_excel_resilient
from aliexpress_orders_sync.models import Order


def test_sync_creates_updates_and_styles_workbook(tmp_path):
    excel_path = tmp_path / "pedidos.xlsx"
    order = Order(
        order_id="1234567890",
        order_date=date(2026, 6, 2),
        item_description="Produto teste",
        quantity=2,
        total_value=Decimal("20.00"),
        currency="USD",
        responsible="Conta A",
        delivery_status="Cancelado",
        tracking_number="LP123456789CN",
        payment_method="Pix",
    )

    created, updated = sync_orders_to_excel([order], excel_path)
    assert (created, updated) == (1, 0)

    changed_order = Order(**{**order.__dict__, "delivery_status": "Entregue"})
    created, updated = sync_orders_to_excel([changed_order], excel_path)
    assert (created, updated) == (0, 1)

    workbook = load_workbook(excel_path, data_only=False)
    sheet = workbook.active

    assert sheet.max_row == 2
    assert sheet["A1"].value == "Data do Pedido"
    assert sheet["B1"].value == "Descrição do Item"
    assert sheet["I1"].value == "Nº Rastreio"
    assert sheet["B2"].value == "Produto teste"
    assert sheet["B2"].alignment.horizontal == "center"
    assert sheet["B2"].alignment.wrap_text is True
    assert sheet["A1"].alignment.horizontal == "center"
    assert sheet["A1"].alignment.wrap_text is True
    assert sheet.column_dimensions["B"].width >= 70
    assert sheet["A2"].alignment.horizontal == "center"
    assert sheet["C2"].alignment.horizontal == "center"
    assert sheet["E2"].value == "=D2/C2"
    assert sheet["F2"].value == "Conta A"
    assert sheet["G2"].value == "Entregue"
    assert sheet["H2"].value == "1234567890"
    assert sheet["I2"].value == "LP123456789CN"
    assert sheet["J1"].value == "Forma de Pagamento"
    assert sheet["J2"].value == "Pix"


def test_sync_normalizes_canceled_status_and_empty_description(tmp_path):
    excel_path = tmp_path / "pedidos.xlsx"
    order = Order(
        order_id="1234567891",
        order_date=date(2026, 6, 2),
        item_description="",
        quantity=1,
        total_value=Decimal("10.00"),
        currency="BRL",
        responsible="Conta B",
        delivery_status="Cancelado",
        tracking_number="",
    )

    sync_orders_to_excel([order], excel_path)
    workbook = load_workbook(excel_path, data_only=False)
    sheet = workbook.active

    assert sheet["B2"].value == "Não identificado"
    assert sheet["F2"].value == "Conta B"
    assert sheet["G2"].value == "Canceled"


def test_sync_updates_same_fixed_workbook_with_account_value(tmp_path):
    excel_path = tmp_path / "pedidos.xlsx"
    first = Order(
        order_id="1234567892",
        order_date=date(2026, 6, 2),
        item_description="Produto",
        quantity=1,
        total_value=Decimal("10.00"),
        currency="BRL",
        responsible="Conta A",
        delivery_status="Awaiting delivery",
        tracking_number="",
    )
    second = Order(**{**first.__dict__, "delivery_status": "Completed"})

    sync_orders_to_excel([first], excel_path)
    sync_orders_to_excel([second], excel_path)

    workbook = load_workbook(excel_path, data_only=False)
    sheet = workbook.active

    assert sheet.max_row == 2
    assert sheet["F2"].value == "Conta A"
    assert sheet["G2"].value == "Completed"


def test_lightweight_update_preserves_existing_detail_columns(tmp_path):
    excel_path = tmp_path / "pedidos.xlsx"
    complete_order = Order(
        order_id="1234567893",
        order_date=date(2026, 6, 2),
        item_description="Produto",
        quantity=1,
        total_value=Decimal("10.00"),
        currency="BRL",
        responsible="Conta A",
        delivery_status="Awaiting delivery",
        tracking_number="LP123456789CN",
        payment_method="Pix",
    )
    lightweight_order = Order(
        **{
            **complete_order.__dict__,
            "delivery_status": "Completed",
            "tracking_number": "",
            "payment_method": "",
        }
    )

    sync_orders_to_excel([complete_order], excel_path)
    sync_orders_to_excel([lightweight_order], excel_path)

    workbook = load_workbook(excel_path, data_only=False)
    sheet = workbook.active

    assert sheet["G2"].value == "Completed"
    assert sheet["I2"].value == "LP123456789CN"
    assert sheet["J2"].value == "Pix"


def test_lightweight_update_preserves_existing_full_description(tmp_path):
    excel_path = tmp_path / "pedidos.xlsx"
    complete_order = Order(
        order_id="1234567894",
        order_date=date(2026, 6, 2),
        item_description="Nome completo do produto com todas as especificações",
        quantity=1,
        total_value=Decimal("10.00"),
        currency="BRL",
        responsible="Conta A",
        delivery_status="Awaiting delivery",
        tracking_number="",
    )
    lightweight_order = Order(
        **{
            **complete_order.__dict__,
            "item_description": "Nome curto",
            "delivery_status": "Completed",
        }
    )

    sync_orders_to_excel([complete_order], excel_path)
    sync_orders_to_excel([lightweight_order], excel_path, preserve_existing_description=True)

    workbook = load_workbook(excel_path, data_only=False)
    sheet = workbook.active

    assert sheet["B2"].value == "Nome completo do produto com todas as especificações"
    assert sheet["G2"].value == "Completed"


def test_invalid_payment_method_is_removed(tmp_path):
    excel_path = tmp_path / "pedidos.xlsx"
    order = Order(
        order_id="1234567895",
        order_date=date(2026, 6, 8),
        item_description="Produto",
        quantity=1,
        total_value=Decimal("62.33"),
        currency="BRL",
        responsible="riquelme",
        delivery_status="Canceled",
        tracking_number="",
        payment_method="Pedido feito em: 8 jun, 2026",
    )

    sync_orders_to_excel([order], excel_path)
    workbook = load_workbook(excel_path, data_only=False)
    sheet = workbook.active

    assert sheet["J2"].value in {"", None}


def test_resilient_sync_queues_orders_when_workbook_is_locked(tmp_path, monkeypatch):
    excel_path = tmp_path / "pedidos.xlsx"
    order = Order(
        order_id="queued-order",
        order_date=date(2026, 6, 8),
        item_description="Produto aguardando",
        quantity=1,
        total_value=Decimal("62.33"),
        currency="BRL",
        responsible="riquelme",
        delivery_status="Canceled",
        tracking_number="",
    )
    original_sync = excel_sync.sync_orders_to_excel

    def locked_sync(*_args, **_kwargs):
        raise PermissionError(13, "Permission denied", str(excel_path))

    monkeypatch.setattr(excel_sync, "sync_orders_to_excel", locked_sync)
    created, updated, queued = sync_orders_to_excel_resilient([order], excel_path)

    assert (created, updated, queued) == (0, 0, 1)
    assert (tmp_path / ".pedidos.pending.json").exists()

    monkeypatch.setattr(excel_sync, "sync_orders_to_excel", original_sync)
    created, updated, queued = sync_orders_to_excel_resilient([], excel_path)

    assert (created, updated, queued) == (1, 0, 0)
    assert not (tmp_path / ".pedidos.pending.json").exists()
    workbook = load_workbook(excel_path, data_only=False)
    assert workbook.active["H2"].value == "queued-order"


def test_sync_sorts_by_date_and_removes_orders_before_2026(tmp_path):
    excel_path = tmp_path / "pedidos.xlsx"
    old_order = Order(
        order_id="old-2025",
        order_date=date(2025, 12, 31),
        item_description="Pedido antigo",
        quantity=1,
        total_value=Decimal("1.00"),
        currency="BRL",
        responsible="neto",
        delivery_status="Completed",
        tracking_number="",
    )
    neto_order = Order(
        order_id="neto-2026-03",
        order_date=date(2026, 3, 18),
        item_description="Pedido neto",
        quantity=1,
        total_value=Decimal("2.00"),
        currency="BRL",
        responsible="neto",
        delivery_status="Completed",
        tracking_number="",
    )
    riquelme_order = Order(
        order_id="riquelme-2026-06",
        order_date=date(2026, 6, 1),
        item_description="Pedido riquelme",
        quantity=1,
        total_value=Decimal("3.00"),
        currency="BRL",
        responsible="riquelme",
        delivery_status="Awaiting delivery",
        tracking_number="",
    )
    neto_same_date = Order(
        order_id="neto-2026-06",
        order_date=date(2026, 6, 1),
        item_description="Pedido neto mesma data",
        quantity=1,
        total_value=Decimal("4.00"),
        currency="BRL",
        responsible="neto",
        delivery_status="Completed",
        tracking_number="",
    )

    sync_orders_to_excel([old_order, neto_order], excel_path)
    sync_orders_to_excel([riquelme_order, neto_same_date], excel_path)

    workbook = load_workbook(excel_path, data_only=False)
    sheet = workbook.active

    order_ids = [sheet.cell(row=row, column=8).value for row in range(2, sheet.max_row + 1)]
    dates = [sheet.cell(row=row, column=1).value.date() for row in range(2, sheet.max_row + 1)]

    assert "old-2025" not in order_ids
    assert dates == [date(2026, 6, 1), date(2026, 6, 1), date(2026, 3, 18)]
    assert set(order_ids[:2]) == {"riquelme-2026-06", "neto-2026-06"}
