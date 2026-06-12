from __future__ import annotations

import json
import re
from dataclasses import asdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Order

HEADERS = [
    "Data do Pedido",
    "Descrição do Item",
    "Quantidade",
    "Valor Total",
    "Valor Unitário",
    "Responsável",
    "Status de Entrega",
    "Número do Pedido",
    "Nº Rastreio",
    "Forma de Pagamento",
]

DESCRIPTION_COLUMN = 2
QUANTITY_COLUMN = 3
TOTAL_COLUMN = 4
UNIT_VALUE_COLUMN = 5
RESPONSIBLE_COLUMN = 6
STATUS_COLUMN = 7
ORDER_ID_COLUMN = 8
TRACKING_COLUMN = 9
PAYMENT_METHOD_COLUMN = 10

HEADER_FILL = PatternFill("solid", fgColor="2F6F57")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SOFT_GRID = Side(style="thin", color="E8EEF2")
SOFT_BORDER = Border(left=SOFT_GRID, right=SOFT_GRID, top=SOFT_GRID, bottom=SOFT_GRID)
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
ZEBRA_FILL = PatternFill("solid", fgColor="F7FAF9")
ORDER_DATE_CUTOFF = date(2026, 1, 1)


def sync_orders_to_excel(
    orders: list[Order],
    excel_path: Path,
    preserve_existing_description: bool = False,
) -> tuple[int, int]:
    """Create or update the workbook, returning (created_rows, updated_rows)."""

    excel_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(excel_path) if excel_path.exists() else _create_workbook()
    sheet = workbook.active

    if not orders:
        if not excel_path.exists():
            _ensure_header(sheet)
            _format_sheet(sheet)
            workbook.save(excel_path)
        return 0, 0

    orders = [order for order in orders if _is_on_or_after_cutoff(order.order_date)]

    _ensure_schema(sheet)
    _remove_orders_before_cutoff(sheet)

    existing_rows = _order_id_to_row(sheet)
    created = 0
    updated = 0

    for order in orders:
        if not order.order_id:
            continue

        if order.order_id in existing_rows:
            row = existing_rows[order.order_id]
            _write_order(sheet, row, order, preserve_existing_description)
            updated += 1
        else:
            row = sheet.max_row + 1
            _write_order(sheet, row, order, False)
            existing_rows[order.order_id] = row
            created += 1

    _remove_orders_before_cutoff(sheet)
    _sanitize_payment_methods(sheet)
    _sort_rows_by_order_date(sheet)
    _format_sheet(sheet)
    _force_formula_recalculation(workbook)
    workbook.save(excel_path)
    return created, updated


def sync_orders_to_excel_resilient(
    orders: list[Order],
    excel_path: Path,
    preserve_existing_description: bool = False,
) -> tuple[int, int, int]:
    """Sync orders or queue them when Excel has locked the workbook."""

    pending_orders = _load_pending_orders(excel_path)
    merged_orders = _merge_orders(pending_orders, orders)
    try:
        created, updated = sync_orders_to_excel(
            merged_orders,
            excel_path,
            preserve_existing_description=preserve_existing_description,
        )
    except PermissionError:
        _save_pending_orders(excel_path, merged_orders)
        return 0, 0, len(merged_orders)

    _remove_pending_file(excel_path)
    return created, updated, 0


def _create_workbook() -> Workbook:
    workbook = Workbook()
    workbook.active.title = "Pedidos"
    return workbook


def _ensure_schema(sheet) -> None:
    current_headers = [sheet.cell(row=1, column=index).value for index in range(1, len(HEADERS) + 1)]
    if current_headers[: len(HEADERS) - 1] == HEADERS[:-1]:
        _ensure_header(sheet)
        return
    if current_headers != HEADERS:
        sheet.delete_rows(1, sheet.max_row)
        sheet.data_validations.dataValidation = []
        sheet.conditional_formatting._cf_rules.clear()
    _ensure_header(sheet)


def _ensure_header(sheet) -> None:
    for column_index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column_index)
        cell.value = header
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = SOFT_BORDER


def _order_id_to_row(sheet) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for row in range(2, sheet.max_row + 1):
        order_id = sheet.cell(row=row, column=ORDER_ID_COLUMN).value
        if order_id:
            mapping[str(order_id)] = row
    return mapping


def _remove_orders_before_cutoff(sheet) -> None:
    for row in range(sheet.max_row, 1, -1):
        order_date = _coerce_date(sheet.cell(row=row, column=1).value)
        if order_date and order_date < ORDER_DATE_CUTOFF:
            sheet.delete_rows(row)


def _sort_rows_by_order_date(sheet) -> None:
    if sheet.max_row <= 2:
        return

    rows = []
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row=row, column=column).value for column in range(1, len(HEADERS) + 1)]
        number_formats = [sheet.cell(row=row, column=column).number_format for column in range(1, len(HEADERS) + 1)]
        if any(value is not None for value in values):
            rows.append((values, number_formats))

    rows.sort(
        key=lambda item: (
            _date_sort_value(item[0][0]),
            str(item[0][RESPONSIBLE_COLUMN - 1] or "").lower(),
            str(item[0][ORDER_ID_COLUMN - 1] or ""),
        ),
        reverse=True,
    )

    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)

    for row_index, (values, number_formats) in enumerate(rows, start=2):
        for column_index, value in enumerate(values, start=1):
            if column_index == UNIT_VALUE_COLUMN:
                value = f"=D{row_index}/C{row_index}"
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.number_format = number_formats[column_index - 1]
        sheet.cell(row=row_index, column=1).number_format = "DD/MM/YYYY"


def _date_sort_value(value) -> int:
    order_date = _coerce_date(value)
    return order_date.toordinal() if order_date else 0


def _is_on_or_after_cutoff(value) -> bool:
    order_date = _coerce_date(value)
    return order_date is None or order_date >= ORDER_DATE_CUTOFF


def _coerce_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, pattern).date()
            except ValueError:
                continue
    return None


def _write_order(sheet, row: int, order: Order, preserve_existing_description: bool = False) -> None:
    total_value = _decimal_to_float(order.total_value)
    quantity = max(1, int(order.quantity or 1))
    existing_description = str(sheet.cell(row=row, column=DESCRIPTION_COLUMN).value or "").strip()
    incoming_description = str(order.item_description or "").strip()
    description = incoming_description or existing_description or "Não identificado"
    if preserve_existing_description and existing_description and existing_description != "Não identificado":
        description = existing_description

    sheet.cell(row=row, column=1, value=order.order_date)
    sheet.cell(row=row, column=DESCRIPTION_COLUMN, value=description)
    sheet.cell(row=row, column=QUANTITY_COLUMN, value=quantity)
    sheet.cell(row=row, column=TOTAL_COLUMN, value=total_value)
    sheet.cell(row=row, column=UNIT_VALUE_COLUMN, value=f"=D{row}/C{row}")
    sheet.cell(row=row, column=RESPONSIBLE_COLUMN, value=order.responsible)
    sheet.cell(row=row, column=STATUS_COLUMN, value=_normalize_status(order.delivery_status))
    sheet.cell(row=row, column=ORDER_ID_COLUMN, value=order.order_id)
    existing_tracking = sheet.cell(row=row, column=TRACKING_COLUMN).value or ""
    existing_payment = _normalize_payment_method(sheet.cell(row=row, column=PAYMENT_METHOD_COLUMN).value)
    sheet.cell(row=row, column=TRACKING_COLUMN, value=order.tracking_number or existing_tracking)
    sheet.cell(
        row=row,
        column=PAYMENT_METHOD_COLUMN,
        value=_normalize_payment_method(order.payment_method) or existing_payment,
    )

    sheet.cell(row=row, column=1).number_format = "DD/MM/YYYY"
    sheet.cell(row=row, column=TOTAL_COLUMN).number_format = _money_format(order.currency)
    sheet.cell(row=row, column=UNIT_VALUE_COLUMN).number_format = _money_format(order.currency)


def _format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    widths = {
        "A": 18,
        "B": 72,
        "C": 14,
        "D": 16,
        "E": 17,
        "F": 24,
        "G": 24,
        "H": 24,
        "I": 28,
        "J": 22,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 28
    _align_cells(sheet)
    _style_rows(sheet)
    last_column = get_column_letter(len(HEADERS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"


def _align_cells(sheet) -> None:
    center = Alignment(horizontal="center", vertical="center")
    description = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=1, max_row=max(sheet.max_row, 1), min_col=1, max_col=len(HEADERS)):
        for cell in row:
            if cell.row == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = description if cell.column == DESCRIPTION_COLUMN else center


def _style_rows(sheet) -> None:
    for row_index in range(2, sheet.max_row + 1):
        row_fill = ZEBRA_FILL if row_index % 2 == 0 else WHITE_FILL
        for column_index in range(1, len(HEADERS) + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.border = SOFT_BORDER
            cell.fill = row_fill


def _force_formula_recalculation(workbook) -> None:
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True


def _normalize_status(status: str) -> str:
    if str(status).strip().lower() in {"cancelado", "canceled", "cancelled", "expired", "expirado"}:
        return "Canceled"
    return status


def _sanitize_payment_methods(sheet) -> None:
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=PAYMENT_METHOD_COLUMN)
        cell.value = _normalize_payment_method(cell.value)


def _normalize_payment_method(value) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if re.search(
        r"\b("
        r"pix|boleto(?:\s+banc[aá]rio)?|"
        r"credit\s*card|debit\s*card|card\s+ending|"
        r"cart[aã]o(?:\s+de)?\s+(?:cr[eé]dito|d[eé]bito)|"
        r"visa|mastercard|master\s*card|american\s+express|amex|elo|hipercard|"
        r"paypal|google\s+pay|apple\s+pay|mercado\s+pago|alipay|webmoney|klarna|"
        r"bank\s+transfer|wire\s+transfer|transfer[eê]ncia\s+banc[aá]ria"
        r")\b",
        text,
        re.I,
    ):
        return text[:80]
    return ""


def _pending_path(excel_path: Path) -> Path:
    return excel_path.with_name(f".{excel_path.stem}.pending.json")


def _load_pending_orders(excel_path: Path) -> list[Order]:
    path = _pending_path(excel_path)
    if not path.exists():
        return []
    try:
        raw_orders = json.loads(path.read_text(encoding="utf-8"))
        return [_order_from_json(item) for item in raw_orders]
    except (OSError, ValueError, TypeError, KeyError):
        return []


def _save_pending_orders(excel_path: Path, orders: list[Order]) -> None:
    path = _pending_path(excel_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    payload = [_order_to_json(order) for order in orders]
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def _remove_pending_file(excel_path: Path) -> None:
    path = _pending_path(excel_path)
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


def _merge_orders(existing: list[Order], incoming: list[Order]) -> list[Order]:
    merged = {order.order_id: order for order in existing if order.order_id}
    for order in incoming:
        if not order.order_id:
            continue
        previous = merged.get(order.order_id)
        if previous:
            order = Order(
                order_id=order.order_id,
                order_date=order.order_date or previous.order_date,
                item_description=_prefer_description(previous.item_description, order.item_description),
                quantity=order.quantity or previous.quantity,
                total_value=order.total_value if order.total_value is not None else previous.total_value,
                currency=order.currency or previous.currency,
                responsible=order.responsible or previous.responsible,
                delivery_status=order.delivery_status or previous.delivery_status,
                tracking_number=order.tracking_number or previous.tracking_number,
                payment_method=_normalize_payment_method(order.payment_method)
                or _normalize_payment_method(previous.payment_method),
            )
        merged[order.order_id] = order
    return list(merged.values())


def _prefer_description(previous: str, incoming: str) -> str:
    previous = str(previous or "").strip()
    incoming = str(incoming or "").strip()
    if not incoming or incoming == "Não identificado":
        return previous or "Não identificado"
    if incoming.endswith("...") and previous and previous != "Não identificado":
        return previous
    return incoming


def _order_to_json(order: Order) -> dict:
    payload = asdict(order)
    payload["order_date"] = order.order_date.isoformat() if order.order_date else None
    payload["total_value"] = str(order.total_value) if order.total_value is not None else None
    return payload


def _order_from_json(payload: dict) -> Order:
    order_date = date.fromisoformat(payload["order_date"]) if payload.get("order_date") else None
    total_value = Decimal(payload["total_value"]) if payload.get("total_value") is not None else None
    return Order(
        order_id=str(payload.get("order_id") or ""),
        order_date=order_date,
        item_description=str(payload.get("item_description") or ""),
        quantity=int(payload.get("quantity") or 1),
        total_value=total_value,
        currency=str(payload.get("currency") or ""),
        responsible=str(payload.get("responsible") or ""),
        delivery_status=str(payload.get("delivery_status") or ""),
        tracking_number=str(payload.get("tracking_number") or ""),
        payment_method=str(payload.get("payment_method") or ""),
    )


def _decimal_to_float(value: Decimal | None) -> float | None:
    if value is None:
        return None
    return float(value)


def _money_format(currency: str) -> str:
    if currency.upper() in {"BRL", "R$", "BR"}:
        return 'R$ #,##0.00'
    if currency.upper() in {"USD", "US$", "$"}:
        return 'US$ #,##0.00'
    return '#,##0.00'
