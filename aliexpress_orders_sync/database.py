from __future__ import annotations

import re
import sqlite3
import threading
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .models import Order

MIN_ORDER_DATE = date(2026, 1, 1)


@dataclass(frozen=True)
class OrderFilters:
    account_key: str = ""
    responsible: str = ""
    status: str = ""
    tracking: str = "all"
    date_from: str = ""
    date_to: str = ""
    search: str = ""
    payment_method: str = ""
    currency: str = ""
    min_total: float | None = None
    max_total: float | None = None


class OrderRepository:
    def __init__(self, database_path: Path) -> None:
        self.database_path = database_path.resolve()
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        self._initialize()

    def upsert_orders(
        self,
        orders: list[Order],
        account_key: str,
        account_email: str = "",
        preserve_existing_description: bool = False,
    ) -> tuple[int, int]:
        created = 0
        updated = 0
        now = datetime.now().isoformat(timespec="seconds")

        with self.lock, self._connect() as connection:
            for order in orders:
                if not order.order_id or (
                    order.order_date is not None and order.order_date < MIN_ORDER_DATE
                ):
                    continue

                existing = connection.execute(
                    """
                    SELECT description, tracking_count, payment_method
                    FROM orders
                    WHERE account_key = ? AND order_id = ?
                    """,
                    (account_key, order.order_id),
                ).fetchone()

                description = str(order.item_description or "").strip() or "Não identificado"
                if (
                    existing
                    and existing["description"]
                    and existing["description"] != "Não identificado"
                    and (
                        preserve_existing_description
                        or description == "Não identificado"
                        or description.endswith("...")
                        or (
                            len(existing["description"]) > len(description)
                            and description.lower() in existing["description"].lower()
                        )
                    )
                ):
                    description = existing["description"]

                tracking_numbers = split_tracking_numbers(order.tracking_number)
                payment_method = normalize_payment_method(order.payment_method)
                if existing and not payment_method:
                    payment_method = existing["payment_method"] or ""

                connection.execute(
                    """
                    INSERT INTO orders (
                        account_key, account_email, order_id, order_date, description,
                        quantity, total_value, currency, responsible, delivery_status,
                        payment_method, tracking_count, first_seen_at, updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(account_key, order_id) DO UPDATE SET
                        account_email = excluded.account_email,
                        order_date = COALESCE(excluded.order_date, orders.order_date),
                        description = excluded.description,
                        quantity = excluded.quantity,
                        total_value = COALESCE(excluded.total_value, orders.total_value),
                        currency = COALESCE(NULLIF(excluded.currency, ''), orders.currency),
                        responsible = COALESCE(NULLIF(excluded.responsible, ''), orders.responsible),
                        delivery_status = COALESCE(NULLIF(excluded.delivery_status, ''), orders.delivery_status),
                        payment_method = COALESCE(NULLIF(excluded.payment_method, ''), orders.payment_method),
                        tracking_count = CASE
                            WHEN excluded.tracking_count > 0 THEN excluded.tracking_count
                            ELSE orders.tracking_count
                        END,
                        updated_at = excluded.updated_at
                    """,
                    (
                        account_key,
                        account_email,
                        order.order_id,
                        order.order_date.isoformat() if order.order_date else None,
                        description,
                        max(1, int(order.quantity or 1)),
                        str(order.total_value) if order.total_value is not None else None,
                        order.currency or "",
                        order.responsible or "",
                        normalize_status(order.delivery_status),
                        payment_method,
                        len(tracking_numbers),
                        now,
                        now,
                    ),
                )

                if tracking_numbers:
                    connection.execute(
                        "DELETE FROM order_tracking WHERE account_key = ? AND order_id = ?",
                        (account_key, order.order_id),
                    )
                    connection.executemany(
                        """
                        INSERT OR IGNORE INTO order_tracking (account_key, order_id, tracking_number)
                        VALUES (?, ?, ?)
                        """,
                        [(account_key, order.order_id, number) for number in tracking_numbers],
                    )

                if existing:
                    updated += 1
                else:
                    created += 1

            connection.commit()
        return created, updated

    def list_orders(
        self,
        filters: OrderFilters | None = None,
        limit: int = 200,
        offset: int = 0,
    ) -> dict:
        filters = filters or OrderFilters()
        where_sql, parameters = self._where(filters)
        limit = max(1, min(limit, 1000))
        offset = max(0, offset)

        with self.lock, self._connect() as connection:
            total = connection.execute(
                f"SELECT COUNT(*) AS total FROM orders o {where_sql}",
                parameters,
            ).fetchone()["total"]
            rows = connection.execute(
                f"""
                SELECT
                    o.*,
                    COALESCE(GROUP_CONCAT(t.tracking_number, ' | '), '') AS tracking_numbers
                FROM orders o
                LEFT JOIN order_tracking t
                    ON t.account_key = o.account_key AND t.order_id = o.order_id
                {where_sql}
                GROUP BY o.account_key, o.order_id
                ORDER BY
                    CASE WHEN o.order_date IS NULL THEN 1 ELSE 0 END,
                    o.order_date DESC,
                    o.updated_at DESC,
                    o.order_id DESC
                LIMIT ? OFFSET ?
                """,
                [*parameters, limit, offset],
            ).fetchall()

        return {
            "total": total,
            "orders": [self._serialize_order(row) for row in rows],
        }

    def dashboard_stats(self, filters: OrderFilters | None = None) -> dict:
        filters = filters or OrderFilters()
        where_sql, parameters = self._where(filters)
        with self.lock, self._connect() as connection:
            row = connection.execute(
                f"""
                SELECT
                    COUNT(*) AS total,
                    SUM(CASE WHEN lower(delivery_status) IN (
                        'completed', 'delivered', 'entregue', 'concluído', 'concluido'
                    ) THEN 1 ELSE 0 END) AS completed,
                    SUM(CASE WHEN lower(delivery_status) IN (
                        'canceled', 'cancelled', 'cancelado', 'expired', 'expirado'
                    ) THEN 1 ELSE 0 END) AS canceled,
                    SUM(CASE WHEN tracking_count > 1 THEN 1 ELSE 0 END) AS multiple_tracking,
                    SUM(CASE WHEN lower(delivery_status) NOT IN (
                        'completed', 'delivered', 'entregue', 'concluído', 'concluido',
                        'canceled', 'cancelled', 'cancelado', 'expired', 'expirado'
                    ) THEN 1 ELSE 0 END) AS in_progress
                FROM orders o
                {where_sql}
                """,
                parameters,
            ).fetchone()

        return {
            "total": int(row["total"] or 0),
            "completed": int(row["completed"] or 0),
            "canceled": int(row["canceled"] or 0),
            "multiple_tracking": int(row["multiple_tracking"] or 0),
            "in_progress": int(row["in_progress"] or 0),
        }

    def available_statuses(self) -> list[str]:
        with self.lock, self._connect() as connection:
            rows = connection.execute(
                """
                SELECT DISTINCT delivery_status
                FROM orders
                WHERE delivery_status IS NOT NULL AND delivery_status != ''
                ORDER BY delivery_status COLLATE NOCASE
                """
            ).fetchall()
        return [row["delivery_status"] for row in rows]

    def filter_options(self) -> dict[str, list[str]]:
        with self.lock, self._connect() as connection:
            payments = connection.execute(
                """
                SELECT DISTINCT payment_method
                FROM orders
                WHERE payment_method IS NOT NULL AND payment_method != ''
                ORDER BY payment_method COLLATE NOCASE
                """
            ).fetchall()
            currencies = connection.execute(
                """
                SELECT DISTINCT currency
                FROM orders
                WHERE currency IS NOT NULL AND currency != ''
                ORDER BY currency
                """
            ).fetchall()
        return {
            "statuses": self.available_statuses(),
            "payment_methods": [row["payment_method"] for row in payments],
            "currencies": [row["currency"] for row in currencies],
        }

    def analytics(self, filters: OrderFilters | None = None) -> dict:
        filters = filters or OrderFilters()
        where_sql, parameters = self._where(filters)
        completed_values = (
            "'completed', 'delivered', 'entregue', 'concluído', 'concluido'"
        )
        canceled_values = (
            "'canceled', 'cancelled', 'cancelado', 'expired', 'expirado'"
        )

        with self.lock, self._connect() as connection:
            summary = connection.execute(
                f"""
                SELECT
                    COUNT(*) AS total_orders,
                    COALESCE(SUM(CAST(total_value AS REAL)), 0) AS gross_value,
                    COALESCE(AVG(CAST(total_value AS REAL)), 0) AS average_ticket,
                    COALESCE(SUM(quantity), 0) AS total_items,
                    SUM(CASE WHEN lower(delivery_status) IN ({completed_values}) THEN 1 ELSE 0 END)
                        AS completed,
                    SUM(CASE WHEN lower(delivery_status) IN ({canceled_values}) THEN 1 ELSE 0 END)
                        AS canceled,
                    SUM(CASE WHEN lower(delivery_status) NOT IN (
                        {completed_values}, {canceled_values}
                    ) THEN 1 ELSE 0 END) AS in_progress,
                    SUM(CASE WHEN tracking_count > 0 THEN 1 ELSE 0 END) AS tracked_orders,
                    SUM(CASE WHEN tracking_count > 1 THEN 1 ELSE 0 END) AS multiple_tracking
                FROM orders o
                {where_sql}
                """,
                parameters,
            ).fetchone()

            monthly = connection.execute(
                f"""
                SELECT
                    substr(order_date, 1, 7) AS month,
                    COUNT(*) AS orders_count,
                    COALESCE(SUM(CAST(total_value AS REAL)), 0) AS total_value
                FROM orders o
                {where_sql}
                GROUP BY substr(order_date, 1, 7)
                HAVING month IS NOT NULL AND month != ''
                ORDER BY month
                """,
                parameters,
            ).fetchall()

            statuses = connection.execute(
                f"""
                SELECT delivery_status AS label, COUNT(*) AS value
                FROM orders o
                {where_sql}
                GROUP BY delivery_status
                ORDER BY value DESC, label COLLATE NOCASE
                """,
                parameters,
            ).fetchall()

            accounts = connection.execute(
                f"""
                SELECT
                    account_key,
                    COALESCE(NULLIF(responsible, ''), account_key) AS label,
                    COUNT(*) AS orders_count,
                    COALESCE(SUM(CAST(total_value AS REAL)), 0) AS total_value
                FROM orders o
                {where_sql}
                GROUP BY account_key, label
                ORDER BY total_value DESC
                """,
                parameters,
            ).fetchall()

            payments = connection.execute(
                f"""
                SELECT
                    COALESCE(NULLIF(payment_method, ''), 'Não informado') AS label,
                    COUNT(*) AS value,
                    COALESCE(SUM(CAST(total_value AS REAL)), 0) AS total_value
                FROM orders o
                {where_sql}
                GROUP BY label
                ORDER BY value DESC, label COLLATE NOCASE
                """,
                parameters,
            ).fetchall()

            currencies = connection.execute(
                f"""
                SELECT
                    COALESCE(NULLIF(currency, ''), 'BRL') AS currency,
                    COUNT(*) AS orders_count,
                    COALESCE(SUM(CAST(total_value AS REAL)), 0) AS total_value
                FROM orders o
                {where_sql}
                GROUP BY currency
                ORDER BY orders_count DESC
                """,
                parameters,
            ).fetchall()

            top_products = connection.execute(
                f"""
                SELECT
                    description AS label,
                    COUNT(*) AS orders_count,
                    COALESCE(SUM(quantity), 0) AS items_count,
                    COALESCE(SUM(CAST(total_value AS REAL)), 0) AS total_value
                FROM orders o
                {where_sql}
                GROUP BY description
                ORDER BY total_value DESC, orders_count DESC
                LIMIT 8
                """,
                parameters,
            ).fetchall()

        return {
            "summary": {
                key: int(summary[key] or 0)
                if key
                in {
                    "total_orders",
                    "total_items",
                    "completed",
                    "canceled",
                    "in_progress",
                    "tracked_orders",
                    "multiple_tracking",
                }
                else round(float(summary[key] or 0), 2)
                for key in summary.keys()
            },
            "monthly": [dict(row) for row in monthly],
            "statuses": [dict(row) for row in statuses],
            "accounts": [dict(row) for row in accounts],
            "payments": [dict(row) for row in payments],
            "currencies": [dict(row) for row in currencies],
            "top_products": [dict(row) for row in top_products],
            "tracking": {
                "none": max(
                    0,
                    int(summary["total_orders"] or 0)
                    - int(summary["tracked_orders"] or 0),
                ),
                "single": max(
                    0,
                    int(summary["tracked_orders"] or 0)
                    - int(summary["multiple_tracking"] or 0),
                ),
                "multiple": int(summary["multiple_tracking"] or 0),
            },
        }

    def count_orders(self) -> int:
        with self.lock, self._connect() as connection:
            return int(connection.execute("SELECT COUNT(*) FROM orders").fetchone()[0])

    def clear_all(self) -> None:
        with self.lock, self._connect() as connection:
            connection.execute("DELETE FROM order_tracking")
            connection.execute("DELETE FROM orders")
            connection.execute("DELETE FROM sync_runs")
            connection.commit()

    def update_account_metadata(self, account_key: str, email: str, responsible: str) -> None:
        with self.lock, self._connect() as connection:
            connection.execute(
                """
                UPDATE orders
                SET account_email = ?, responsible = ?, updated_at = ?
                WHERE account_key = ?
                """,
                (
                    email,
                    responsible,
                    datetime.now().isoformat(timespec="seconds"),
                    account_key,
                ),
            )
            connection.commit()

    def import_legacy_workbook(self, excel_path: Path, account_keys_by_responsible: dict[str, str]) -> int:
        if self.count_orders() or not excel_path.exists():
            return 0

        workbook = load_workbook(excel_path, data_only=True, read_only=True)
        sheet = workbook.active
        grouped: dict[tuple[str, str], list[Order]] = {}
        try:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if len(row) < 8 or not row[7]:
                    continue
                responsible = str(row[5] or "").strip()
                account_key = account_keys_by_responsible.get(responsible.lower(), responsible or "importado")
                order = Order(
                    order_id=str(row[7]),
                    order_date=coerce_date(row[0]),
                    item_description=str(row[1] or "Não identificado"),
                    quantity=max(1, int(row[2] or 1)),
                    total_value=coerce_decimal(row[3]),
                    currency="BRL",
                    responsible=responsible,
                    delivery_status=str(row[6] or "Não identificado"),
                    tracking_number=str(row[8] or "") if len(row) > 8 else "",
                    payment_method=str(row[9] or "") if len(row) > 9 else "",
                )
                grouped.setdefault((account_key, responsible), []).append(order)
        finally:
            workbook.close()

        imported = 0
        for (account_key, _responsible), orders in grouped.items():
            created, _updated = self.upsert_orders(orders, account_key)
            imported += created
        return imported

    def record_sync(
        self,
        account_key: str,
        mode: str,
        read_count: int,
        created_count: int,
        updated_count: int,
        status: str = "success",
        error_message: str = "",
    ) -> None:
        with self.lock, self._connect() as connection:
            connection.execute(
                """
                INSERT INTO sync_runs (
                    account_key, mode, read_count, created_count, updated_count,
                    status, error_message, finished_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    account_key,
                    mode,
                    read_count,
                    created_count,
                    updated_count,
                    status,
                    error_message,
                    datetime.now().isoformat(timespec="seconds"),
                ),
            )
            connection.commit()

    def _initialize(self) -> None:
        with self._connect() as connection:
            connection.executescript(
                """
                PRAGMA journal_mode = WAL;
                PRAGMA foreign_keys = ON;

                CREATE TABLE IF NOT EXISTS orders (
                    account_key TEXT NOT NULL,
                    account_email TEXT NOT NULL DEFAULT '',
                    order_id TEXT NOT NULL,
                    order_date TEXT,
                    description TEXT NOT NULL DEFAULT 'Não identificado',
                    quantity INTEGER NOT NULL DEFAULT 1,
                    total_value TEXT,
                    currency TEXT NOT NULL DEFAULT '',
                    responsible TEXT NOT NULL DEFAULT '',
                    delivery_status TEXT NOT NULL DEFAULT 'Não identificado',
                    payment_method TEXT NOT NULL DEFAULT '',
                    tracking_count INTEGER NOT NULL DEFAULT 0,
                    first_seen_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY (account_key, order_id)
                );

                CREATE TABLE IF NOT EXISTS order_tracking (
                    account_key TEXT NOT NULL,
                    order_id TEXT NOT NULL,
                    tracking_number TEXT NOT NULL,
                    PRIMARY KEY (account_key, order_id, tracking_number),
                    FOREIGN KEY (account_key, order_id)
                        REFERENCES orders(account_key, order_id)
                        ON DELETE CASCADE
                );

                CREATE TABLE IF NOT EXISTS sync_runs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    account_key TEXT NOT NULL,
                    mode TEXT NOT NULL,
                    read_count INTEGER NOT NULL DEFAULT 0,
                    created_count INTEGER NOT NULL DEFAULT 0,
                    updated_count INTEGER NOT NULL DEFAULT 0,
                    status TEXT NOT NULL,
                    error_message TEXT NOT NULL DEFAULT '',
                    finished_at TEXT NOT NULL
                );

                CREATE INDEX IF NOT EXISTS idx_orders_date ON orders(order_date DESC);
                CREATE INDEX IF NOT EXISTS idx_orders_status ON orders(delivery_status);
                CREATE INDEX IF NOT EXISTS idx_orders_account ON orders(account_key);
                CREATE INDEX IF NOT EXISTS idx_orders_tracking_count ON orders(tracking_count);

                DELETE FROM orders
                WHERE order_date IS NOT NULL AND order_date < '2026-01-01';
                """
            )
            connection.commit()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database_path, timeout=30)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA busy_timeout = 30000")
        connection.execute("PRAGMA foreign_keys = ON")
        return connection

    def _where(self, filters: OrderFilters) -> tuple[str, list]:
        clauses: list[str] = []
        parameters: list = []

        if filters.account_key:
            clauses.append("o.account_key = ?")
            parameters.append(filters.account_key)
        if filters.responsible:
            clauses.append("lower(o.responsible) = lower(?)")
            parameters.append(filters.responsible)
        if filters.status:
            clauses.append("lower(o.delivery_status) = lower(?)")
            parameters.append(filters.status)
        if filters.date_from:
            clauses.append("o.order_date >= ?")
            parameters.append(filters.date_from)
        if filters.date_to:
            clauses.append("o.order_date <= ?")
            parameters.append(filters.date_to)
        if filters.payment_method:
            clauses.append("lower(o.payment_method) = lower(?)")
            parameters.append(filters.payment_method)
        if filters.currency:
            clauses.append("upper(o.currency) = upper(?)")
            parameters.append(filters.currency)
        if filters.min_total is not None:
            clauses.append("CAST(o.total_value AS REAL) >= ?")
            parameters.append(filters.min_total)
        if filters.max_total is not None:
            clauses.append("CAST(o.total_value AS REAL) <= ?")
            parameters.append(filters.max_total)
        if filters.tracking == "multiple":
            clauses.append("o.tracking_count > 1")
        elif filters.tracking == "single":
            clauses.append("o.tracking_count = 1")
        elif filters.tracking == "none":
            clauses.append("o.tracking_count = 0")
        if filters.search:
            search = f"%{filters.search.strip()}%"
            clauses.append(
                """
                (
                    o.description LIKE ? OR o.order_id LIKE ? OR o.responsible LIKE ?
                    OR o.payment_method LIKE ? OR EXISTS (
                        SELECT 1 FROM order_tracking st
                        WHERE st.account_key = o.account_key
                          AND st.order_id = o.order_id
                          AND st.tracking_number LIKE ?
                    )
                )
                """
            )
            parameters.extend([search, search, search, search, search])

        return (f"WHERE {' AND '.join(clauses)}" if clauses else ""), parameters

    @staticmethod
    def _serialize_order(row: sqlite3.Row) -> dict:
        total = float(row["total_value"]) if row["total_value"] not in {None, ""} else None
        quantity = max(1, int(row["quantity"] or 1))
        return {
            "account_key": row["account_key"],
            "account_email": row["account_email"],
            "order_id": row["order_id"],
            "order_date": display_date(row["order_date"]),
            "order_date_iso": row["order_date"] or "",
            "description": row["description"],
            "quantity": quantity,
            "total": total,
            "unit_value": round(total / quantity, 2) if total is not None else None,
            "currency": row["currency"],
            "responsible": row["responsible"],
            "status": row["delivery_status"],
            "payment_method": row["payment_method"],
            "tracking_numbers": split_tracking_numbers(row["tracking_numbers"]),
            "tracking_count": int(row["tracking_count"] or 0),
            "updated_at": row["updated_at"],
        }


def split_tracking_numbers(value: str) -> list[str]:
    parts = re.split(r"\s*(?:\||;|,|\n)\s*", str(value or ""))
    result: list[str] = []
    for part in parts:
        cleaned = part.strip()
        if len(cleaned) < 6 or cleaned in result:
            continue
        result.append(cleaned)
    return result


def normalize_status(status: str) -> str:
    value = str(status or "Não identificado").strip()
    if value.lower() in {"cancelado", "canceled", "cancelled", "expired", "expirado"}:
        return "Canceled"
    return value


def normalize_payment_method(value: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if re.search(
        r"\b("
        r"pix|boleto(?:\s+banc[aá]rio)?|"
        r"credit\s*card|debit\s*card|card\s+ending|"
        r"cart[aã]o(?:\s+de)?\s+(?:cr[eé]dito|d[eé]bito)|"
        r"visa|mastercard|master\s*card|american\s+express|amex|elo|hipercard|"
        r"paypal|google\s+pay|apple\s+pay|mercado\s+pago|alipay|webmoney|klarna|"
        r"bank\s+transfer|wire\s+transfer|transfer[eê]ncia\s+banc[aá]ria"
        r")\b",
        text,
        re.I,
    ):
        return text[:80]
    return ""


def display_date(value: str | None) -> str:
    if not value:
        return ""
    try:
        return date.fromisoformat(value).strftime("%d/%m/%Y")
    except ValueError:
        return value


def coerce_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(value.strip(), pattern).date()
            except ValueError:
                continue
    return None


def coerce_decimal(value) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value))
    except Exception:
        return None
